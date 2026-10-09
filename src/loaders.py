"""Format-specific document loaders (FR-B2, FR-B3)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".md", ".pdf", ".docx"}


def load_markdown(path: Path) -> str:
    """Read a markdown file as UTF-8 text."""
    return path.read_text(encoding="utf-8")


def load_pdf(path: Path) -> str:
    """Extract text from a PDF using pypdf.

    pypdf is used deliberately here, not pdfplumber/pdfminer. On this corpus,
    pdfplumber and pdfminer.six both extract the Arabic PDFs with every line
    character-reversed (a classic RTL/bidi extraction bug — exactly the
    "scrambled glyphs" symptom that's a red flag for bad extraction). pypdf
    handles the same files correctly out of the box. Verified by hand on all
    9 PDFs (6 English, 3 Arabic) before this choice was locked in — re-check
    this if you ever swap PDF libraries.
    """
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def load_docx(path: Path) -> str:
    """Extract paragraph and table-cell text from a .docx file."""
    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


LOADERS = {".md": load_markdown, ".pdf": load_pdf, ".docx": load_docx}


def load_document(path: Path) -> str:
    """Load a single document, failing loudly on any unsupported format (FR-B2).

    Args:
        path: Path to a .md, .pdf, or .docx file.

    Returns:
        The extracted raw text.

    Raises:
        ValueError: if the extension is unsupported, or extraction yields
            only whitespace (a strong signal the loader/encoding is wrong).
    """
    ext = path.suffix.lower()
    if ext not in LOADERS:
        raise ValueError(
            f"Unsupported file format '{ext}' for '{path.name}'. "
            f"Supported formats: {sorted(SUPPORTED_EXTENSIONS)}."
        )
    text = LOADERS[ext](path)
    if not text.strip():
        raise ValueError(
            f"Extracted empty text from '{path.name}' — check the loader/encoding."
        )
    return text


def strip_markdown_headings(text: str) -> str:
    """Drop leading '#' markers from markdown headings, keeping the heading text."""
    return re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
